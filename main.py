from datetime import datetime
from openpyxl import Workbook

wb = Workbook()

ws = wb.active
ws.title = "新闻"

ws['A1'] = "日期"
ws['B1'] = "标题"
ws['C1'] = "来源"

today = datetime.now().strftime("%Y-%m-%d")

ws.append([today,
           "测试新闻",
           "GitHub Actions"])

filename = f"日报_{today}.xlsx"

wb.save(filename)

print("Excel创建成功：", filename)
